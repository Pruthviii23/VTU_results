import os
import logging

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import config
import sgpa as sgpa_calc

log = logging.getLogger(__name__)

# ============================================================
# STYLE CONSTANTS
# ============================================================

_HEADER_FILL   = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
_SUBHEAD_FILL  = PatternFill("solid", start_color="2E75B6", end_color="2E75B6")
_ALT_FILL      = PatternFill("solid", start_color="D6E4F0", end_color="D6E4F0")
_FAIL_FILL     = PatternFill("solid", start_color="FFD7D7", end_color="FFD7D7")  # red   — Fail
_ABSENT_FILL   = PatternFill("solid", start_color="FFF3CD", end_color="FFF3CD")  # amber — Absent
_WITHHELD_FILL = PatternFill("solid", start_color="E8D5F5", end_color="E8D5F5")  # purple — Withheld
_NE_FILL       = PatternFill("solid", start_color="D1D5DB", end_color="D1D5DB")  # grey  — Not Eligible
_WHITE_FILL    = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")

_HEADER_FONT   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
_SUBHEAD_FONT  = Font(name="Arial", bold=True, color="FFFFFF", size=9)
_DATA_FONT     = Font(name="Arial", size=9)
_BOLD_FONT     = Font(name="Arial", bold=True, size=9)

_CENTER        = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT          = Alignment(horizontal="left",   vertical="center", wrap_text=True)

_THIN          = Side(style="thin", color="AAAAAA")
_BORDER        = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_SUB_FIELDS    = ["Internal", "External", "Total", "Result"]


# ============================================================
# INTERNAL HELPERS
# ============================================================

def _style_cell(cell, font=None, fill=None, alignment=None, border=None):
    if font:      cell.font      = font
    if fill:      cell.fill      = fill
    if alignment: cell.alignment = alignment
    if border:    cell.border    = border


def _collect_subject_codes(rows: list[dict]) -> list[str]:
    """
    Given all student rows for one semester, returns a sorted,
    deduplicated list of subject codes found across all students.
    This handles elective differences — everyone's subjects are unioned.
    """
    codes = set()
    for student in rows:
        for subj in student.get("subjects", []):
            code = subj["code"].strip().upper()
            if code:
                codes.add(code)
    return sorted(codes)


def _build_subject_index(subjects: list[dict]) -> dict:
    """Map subject code → subject dict for fast lookup."""
    return {s["code"].strip().upper(): s for s in subjects}


def _write_sheet(ws, sem_num: int, rows: list[dict], scheme: str = "2022"):
    """
    Write one semester's worth of data to `ws`.

    Row 1 : Group headers  (USN | Name | <SubjectCode> spanning 4 cols | ...)
    Row 2 : Sub-headers    (     |      | Internal | External | Total | Result | ...)
    Row 3+: Student data
    """
    all_codes = _collect_subject_codes(rows)

    if not all_codes:
        log.warning(f"Sem {sem_num}: no subject codes found — sheet will be empty.")

    # ------------------------------------------------
    # Build column map: col_index → (header, sub_header)
    # col 1 = USN, col 2 = Name, then groups of 4 per subject
    # ------------------------------------------------
    col = 1
    fixed_cols  = {1: "USN", 2: "Name"}
    subject_col = {}    # code → starting column index

    col = 3
    for code in all_codes:
        subject_col[code] = col
        col += 4          # Internal, External, Total, Result

    # SGPA columns come after all subjects
    sgpa_col         = col        # SGPA value
    credits_col      = col + 1    # Total credits
    earned_col       = col + 2    # Earned credits

    # ------------------------------------------------
    # Row 1: headers
    # ------------------------------------------------
    ws.cell(row=1, column=1, value="USN").alignment  = _CENTER
    ws.cell(row=1, column=2, value="Name").alignment = _CENTER
    ws.merge_cells(start_row=1, end_row=2, start_column=1, end_column=1)
    ws.merge_cells(start_row=1, end_row=2, start_column=2, end_column=2)

    for cell in [ws.cell(row=1, column=1), ws.cell(row=1, column=2),
                 ws.cell(row=2, column=1), ws.cell(row=2, column=2)]:
        _style_cell(cell, font=_HEADER_FONT, fill=_HEADER_FILL,
                    alignment=_CENTER, border=_BORDER)

    for code, start_col in subject_col.items():
        end_col = start_col + 3
        ws.merge_cells(
            start_row=1, end_row=1,
            start_column=start_col, end_column=end_col
        )
        cell = ws.cell(row=1, column=start_col, value=code)
        _style_cell(cell, font=_HEADER_FONT, fill=_HEADER_FILL,
                    alignment=_CENTER, border=_BORDER)

        for i, field in enumerate(_SUB_FIELDS):
            sub_cell = ws.cell(row=2, column=start_col + i, value=field)
            _style_cell(sub_cell, font=_SUBHEAD_FONT, fill=_SUBHEAD_FILL,
                        alignment=_CENTER, border=_BORDER)

    # SGPA header group
    _SGPA_FILL = PatternFill("solid", start_color="1D6B3E", end_color="1D6B3E")
    ws.merge_cells(start_row=1, end_row=1,
                   start_column=sgpa_col, end_column=earned_col)
    grp = ws.cell(row=1, column=sgpa_col, value="SGPA")
    _style_cell(grp, font=_HEADER_FONT, fill=_SGPA_FILL,
                alignment=_CENTER, border=_BORDER)

    for c, label in [(sgpa_col, "SGPA"), (credits_col, "Total Credits"), (earned_col, "Earned Credits")]:
        cell = ws.cell(row=2, column=c, value=label)
        _style_cell(cell, font=_SUBHEAD_FONT, fill=_SGPA_FILL,
                    alignment=_CENTER, border=_BORDER)

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18

    # ------------------------------------------------
    # Row 3+: student data
    # ------------------------------------------------
    for r_idx, student in enumerate(rows, start=3):
        alt = (r_idx % 2 == 0)
        bg  = _ALT_FILL if alt else _WHITE_FILL

        def write(c, value, bold=False, fail=False, absent=False,
                  withheld=False, ne=False):
            cell = ws.cell(row=r_idx, column=c, value=value)
            if withheld:
                fill = _WITHHELD_FILL
            elif ne:
                fill = _NE_FILL
            elif absent:
                fill = _ABSENT_FILL
            elif fail:
                fill = _FAIL_FILL
            else:
                fill = bg
            _style_cell(
                cell,
                font      = _BOLD_FONT if bold else _DATA_FONT,
                fill      = fill,
                alignment = _LEFT if c == 2 else _CENTER,
                border    = _BORDER
            )

        write(1, student.get("USN", ""),  bold=True)
        write(2, student.get("Name", ""), bold=True)

        subj_index = _build_subject_index(student.get("subjects", []))

        for code, start_col in subject_col.items():
            subj = subj_index.get(code)

            if subj is None:
                # Student didn't take this elective — leave blank
                for i in range(4):
                    cell = ws.cell(row=r_idx, column=start_col + i, value="")
                    _style_cell(cell, font=_DATA_FONT, fill=bg,
                                alignment=_CENTER, border=_BORDER)
            else:
                result  = subj.get("result", "")
                r_upper = result.upper()
                failed   = r_upper in ("F", "--")
                absent   = r_upper in ("A", "AB", "ABSENT")
                withheld = r_upper in ("W", "WITHHELD")
                ne       = r_upper in ("X", "NE", "NOT ELIGIBLE")

                write(start_col,     subj.get("internal", ""), fail=failed, absent=absent, withheld=withheld, ne=ne)
                write(start_col + 1, subj.get("external", ""), fail=failed, absent=absent, withheld=withheld, ne=ne)
                write(start_col + 2, subj.get("total",    ""), fail=failed, absent=absent, withheld=withheld, ne=ne)
                write(start_col + 3, result,                    fail=failed, absent=absent, withheld=withheld, ne=ne)

        # SGPA columns
        sgpa_result = sgpa_calc.compute_sgpa(student.get("subjects", []), scheme, student.get("USN", ""))
        sgpa_val    = sgpa_result["sgpa"]

        _SGPA_DATA_FILL = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
        _SGPA_DATA_FONT = Font(name="Arial", bold=True, size=9, color="1D6B3E")

        for c, val in [
            (sgpa_col,    f"{sgpa_val:.2f}" if sgpa_val is not None else "—"),
            (credits_col, sgpa_result["total_credits"]),
            (earned_col,  sgpa_result["earned_credits"]),
        ]:
            cell = ws.cell(row=r_idx, column=c, value=val)
            _style_cell(cell, font=_SGPA_DATA_FONT, fill=_SGPA_DATA_FILL,
                        alignment=_CENTER, border=_BORDER)

    # ------------------------------------------------
    # Column widths
    # ------------------------------------------------
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 28
    for code, start_col in subject_col.items():
        for i in range(4):
            ltr = get_column_letter(start_col + i)
            ws.column_dimensions[ltr].width = 11
    for c in [sgpa_col, credits_col, earned_col]:
        ws.column_dimensions[get_column_letter(c)].width = 14

    # Freeze header rows
    ws.freeze_panes = ws.cell(row=3, column=3)


def _write_summary_sheet(ws, all_students: list[dict], target_sem: int):
    """
    Summary sheet: one row per student, columns = USN | Name | Pass | Fail | Backlogs
    Only counts subjects from the target semester.
    """
    headers = ["USN", "Name", f"Sem {target_sem} Pass", f"Sem {target_sem} Fail", "Has Backlog"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        _style_cell(cell, font=_HEADER_FONT, fill=_HEADER_FILL,
                    alignment=_CENTER, border=_BORDER)
    ws.row_dimensions[1].height = 22

    for r, student in enumerate(all_students, start=2):
        sem_subjects = student["semesters"].get(target_sem, [])
        passes   = sum(1 for s in sem_subjects if s.get("result", "").upper() == "P")
        fails    = sum(1 for s in sem_subjects if s.get("result", "").upper() == "F")
        backlog  = len(student["semesters"]) > 1   # has data from other semesters too

        alt = (r % 2 == 0)
        bg  = _ALT_FILL if alt else _WHITE_FILL

        for c, val in enumerate(
            [student["USN"], student["Name"], passes, fails, "Yes" if backlog else "No"], 1
        ):
            cell = ws.cell(row=r, column=c, value=val)
            _style_cell(cell, font=_DATA_FONT, fill=bg,
                        alignment=_CENTER if c != 2 else _LEFT, border=_BORDER)

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 28
    for ltr in ["C", "D", "E"]:
        ws.column_dimensions[ltr].width = 14
    ws.freeze_panes = ws.cell(row=2, column=3)


# ============================================================
# PUBLIC API
# ============================================================

def write_excel(all_students: list[dict], output_path: str,
                target_semester: int, scheme: str = "2022"):
    """
    Writes one Excel workbook with:
      - Sheet per semester (target + any backlog semesters found)
      - Summary sheet
    Columns are built dynamically from whatever subject codes appear.
    """
    if not all_students:
        log.warning("No student data to write.")
        return

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    from collections import defaultdict
    semester_rows: dict[int, list[dict]] = defaultdict(list)

    for student in all_students:
        for sem_num, subjects in student["semesters"].items():
            semester_rows[sem_num].append({
                "USN":      student["USN"],
                "Name":     student["Name"],
                "subjects": subjects,
            })

    wb = Workbook()
    wb.remove(wb.active)

    ordered_sems = sorted(
        semester_rows.keys(),
        key=lambda s: (s != target_semester, s)
    )

    for sem in ordered_sems:
        ws = wb.create_sheet(title=f"Sem {sem}")
        _write_sheet(ws, sem, semester_rows[sem], scheme)
        log.info(f"Wrote sheet 'Sem {sem}' with {len(semester_rows[sem])} rows.")

    ws_summary = wb.create_sheet(title="Summary")
    _write_summary_sheet(ws_summary, all_students, target_semester)

    try:
        wb.save(output_path)
        log.info(f"Excel saved: {output_path}")
    except PermissionError:
        fallback = output_path.replace(".xlsx", "_backup.xlsx")
        wb.save(fallback)
        log.error(
            f"Output file was open — saved to fallback: {fallback}"
        )